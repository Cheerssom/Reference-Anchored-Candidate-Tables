"""Read-only verification of the frozen manuscript, sources and attachments."""
from pathlib import Path
import hashlib,json,zipfile
import fitz
from docx import Document

H=Path(__file__).resolve().parent
manifest=json.loads((H/'submission_freeze_manifest.json').read_text(encoding='utf-8'))
checked=0
for name,sha in manifest['file_sha256'].items():
    p=H/name
    # Delivery ZIPs are kept locally; all published sources/PDFs are required.
    if not p.exists() and name in manifest['local_only_archives']:
        continue
    assert p.exists(),name
    assert hashlib.sha256(p.read_bytes()).hexdigest()==sha,name
    checked+=1
for name in manifest['local_only_archives']:
    p=H/name
    if not p.exists():continue
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None,name
        hashes=json.loads(z.read('file_hashes.json'))
        for item,sha in hashes.items():
            assert hashlib.sha256(z.read(item)).hexdigest()==sha,(name,item)
title=manifest['title']
for name in ['main','supplementary']:
    with fitz.open(H/f'{name}.pdf') as pdf:
        assert len(pdf)==manifest['pages'][name]
        text=' '.join(pdf[0].get_text().split())
        assert title in text,(name,'title')
        if name=='main':
            assert 'matched HyCal transfer' in text
            assert 'ten partitions per encoder.' in text and 'Keywords:' in text
for name in ['title_page.docx','cover_letter.docx']:
    text='\n'.join(p.text for p in Document(H/name).paragraphs)
    assert title in text and 'Junyi Zhang' in text and 'cheerssom0876@gmail.com' in text
high=[p.text for p in Document(H/'Highlights.docx').paragraphs][1:]
assert list(map(len,high))==manifest['highlight_characters']
main=(H/'main.tex').read_text(encoding='utf-8')
supp=(H/'supplementary.tex').read_text(encoding='utf-8')
assert 'If both accuracies match within every class' in supp
assert 'Matching overall macro accuracies alone does not imply' in supp
assert main.index('Learn then Test framework')<main.index(r'\begin{theorem}[Macro protection')
assert main.index('We adapt bounded-mean betting')<main.index(r'\begin{theorem}[Direct paired accuracy test]')
assert r'\subsection{Query-adaptive alignment improves base--novel decisions}' in main
assert r'\subsection{Query-adaptive alignment expands retention frontiers}' not in main
assert 'L/14 retained readout reaches' in main
for name in ['table_asymmetric_recognition.tex','table_asymmetric_centering.tex','table_asymmetric_admission.tex']:
    table=(H/name).read_text(encoding='utf-8')
    assert r'\textnormal{CLIP ViT-L/14}' in table
    assert r'\textnormal{CLIP ViT-B/32}' in table
with fitz.open(H/'framework.pdf') as figure:
    figure_text=' '.join(p.get_text() for p in figure)
    assert 'Asymmetric readout' in figure_text and 'readout evidence' in figure_text
    assert 'Candidate table' not in figure_text and 'candidate evidence' not in figure_text
print(json.dumps({'complete':True,'frozen_files_checked':checked,
                  'freeze_id':manifest['freeze_id'],'pages':manifest['pages'],
                  'abstract_words':manifest['abstract_words'],
                  'journal_specific_rules_directly_verified':False},indent=2))

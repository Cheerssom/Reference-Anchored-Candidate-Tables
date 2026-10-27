"""Prepare editable KBS submission attachments and reproducible source archive."""
from pathlib import Path
import hashlib,json,re,zipfile
from docx import Document
from docx.shared import Inches,Pt

H=Path(__file__).resolve().parent
title='Reference-anchored asymmetric class expansion with query-adaptive evidence alignment'
highlights=[
 'Fixed base references align asymmetric evidence without vocabulary-induced drift.',
 'Query-adaptive margins improve base--novel recognition without extra routers.',
 'Matched HyCal gains reach 10.93 and 13.74 points on ImageNet-R.',
 'Independent calibration selects admission under a base-accuracy requirement.',
 'Target calibration transfers frozen classifiers and prototypes to new renditions.'
]
assert 3<=len(highlights)<=5 and max(map(len,highlights))<=85
(H/'Highlights.txt').write_text('\n'.join('- '+v for v in highlights)+'\n',encoding='utf-8')
letter=f"""Dear Editors of Knowledge-Based Systems,

Please consider our manuscript, "{title}," as an original Research Article.

The paper addresses class expansion when a trained base classifier and arriving few-shot classes provide asymmetric evidence. Fixed base references align learned scores, semantic anchors and visual prototypes while preserving existing scores during vocabulary growth. Query normalization induces an explicit sample-adaptive admission margin without an auxiliary routing network. Independent calibration selects the global offset under a specified base-accuracy requirement, enabling retention-constrained expansion and frozen-model transfer.

Experiments cover three source datasets, ten class partitions, two frozen CLIP encoders and five arrivals. In matched transfers on ImageNet-R, average incremental accuracy exceeds HyCal by 10.93 and 13.74 percentage points. A fixed-fusion centering ablation improves CIFAR L/14 novel accuracy by 41.73 points after independent offset recalibration; separately optimized normalization schemes confirm ImageNet-R capacity gains of 3.60 and 5.29 points. Under a common direct accuracy certificate, the original retained ImageNet-R L/14 readout reaches 81.61% novel accuracy, versus 76.76% for dense prototypes, with both admitted in all ten partitions. Frozen source models and prototypes reach 86.82% and 58.96% novel accuracy on ImageNet-Sketch after target calibration, with all evaluated empirical retention targets met. The contribution connects a concrete evidence-alignment mechanism to sustained recognition and useful admission. Code and experimental records accompany the manuscript.

The manuscript is original, has not been formally published and is not under consideration elsewhere. I approve submission to Knowledge-Based Systems, declare no competing interests, and received no external funding. AI assistance is disclosed. Code and source-hashed records are available at https://github.com/Cheerssom/Reference-Anchored-Candidate-Tables .

Sincerely,
Junyi Zhang
Shandong University, China
cheerssom0876@gmail.com
"""
(H/'cover_letter.md').write_text(letter,encoding='utf-8')

def doc(name,heading,paragraphs,bullets=False):
 d=Document();section=d.sections[0]
 section.top_margin=section.bottom_margin=Inches(.8)
 section.left_margin=section.right_margin=Inches(.85)
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
 normal.paragraph_format.space_after=Pt(7)
 if heading:d.add_heading(heading,level=1)
 for text in paragraphs:d.add_paragraph(text,style='List Bullet' if bullets else None)
 d.save(H/name)

doc('Highlights.docx','Highlights',highlights,True)
doc('cover_letter.docx','',letter.strip().split('\n\n'))
doc('title_page.docx',title,['Junyi Zhang','Shandong University, China','Corresponding author: Junyi Zhang','Email: cheerssom0876@gmail.com'])
doc('declarations.docx','Declarations',[
'CRediT authorship contribution statement: Junyi Zhang: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing - original draft, Writing - review and editing.',
'Funding: This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.',
'Declaration of competing interest: The author declares that there are no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.',
'Data availability: Public datasets and the CLIP checkpoint are obtained separately. Code and experimental records: https://github.com/Cheerssom/Reference-Anchored-Candidate-Tables .',
'AI assistance: OpenAI Codex assisted literature discovery, manuscript drafting and revision, experiment and verification scripts, and LaTeX and table preparation. The author is responsible for the accuracy, interpretation and originality of the submitted work and for reviewing and approving the final manuscript and supporting materials.'
])

sources=[H/n for n in ['main.tex','main.bbl','supplementary.tex','supplementary.bbl','references.bib','elsarticle.cls','elsarticle-num.bst','numbers.tex','retention_numbers.tex','strengthening_results.tex','depth_results.tex','paired_accuracy_proof.tex'] if (H/n).exists()]
sources += [H/n for n in ['elsarticle.dtx','elsarticle.ins','template_README','template_manifest.txt','elsarticle-template-num.tex','template_provenance.json']]
sources+=list(H.glob('table_*.tex'))
sources+=[H/'validation_budget_results.tex']
sources+=[p for p in H.glob('*.pdf') if p.name!='logo.pdf']
sources=[p for p in sources if p.name not in ['main.pdf','supplementary.pdf','review_complete.pdf','cover_letter.pdf','Highlights.pdf','title_page.pdf','declarations.pdf']]
sources+=list(H.glob('*.docx'))+[H/'Highlights.txt']
sources=sorted(set(sources))
manifest={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
with zipfile.ZipFile(H/'candidate_tables_kbs_source.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in sources:z.write(p,p.name)
 z.writestr('file_hashes.json',json.dumps(manifest,indent=2))
with zipfile.ZipFile(H/'candidate_tables_kbs_source.zip') as z:
 assert z.testzip() is None
 assert all('/' not in n for n in z.namelist())
with zipfile.ZipFile(H/'candidate_tables_kbs_supplement_source.zip','w',zipfile.ZIP_DEFLATED) as z:
 supplement_manifest={}
 for p in sources:
  if p.name not in ['main.tex','main.bbl']:
   z.write(p,p.name);supplement_manifest[p.name]=hashlib.sha256(p.read_bytes()).hexdigest()
 z.write(H/'main.aux','main.aux')
 supplement_manifest['main.aux']=hashlib.sha256((H/'main.aux').read_bytes()).hexdigest()
 z.writestr('BUILD.txt','Compile supplementary.tex. main.aux supplies references to the main manuscript.\n')
 z.writestr('file_hashes.json',json.dumps(supplement_manifest,indent=2))
print(json.dumps({'highlight_characters':list(map(len,highlights)),'flat_source_files':len(sources)},indent=2))
